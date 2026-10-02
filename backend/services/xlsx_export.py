"""Turn list of transaction dicts into an Excel (.xlsx) workbook."""
import io
import re
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from services.csv_export import export_fields

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HEADERS = {"date": "Date", "description": "Description", "amount": "Amount", "type": "Type", "category": "Category"}
_WIDTHS = {"date": 12, "description": 48, "amount": 14, "type": 10, "category": 20}
_NUMBER_RE = re.compile(r"^-?\d+(\.\d+)?$")


def _amount_value(raw: Any) -> float | str:
    """Real number when the amount parses (so Excel can sum it), else the original text."""
    text = str(raw or "").replace(",", "").strip()
    return float(text) if _NUMBER_RE.match(text) else str(raw or "")


def transactions_to_xlsx(transactions: list[dict[str, Any]], include_categories: bool = True) -> bytes:
    """Return .xlsx bytes: bold frozen header, numeric amounts, sized columns, autofilter."""
    fields = export_fields(include_categories)
    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"

    ws.append([_HEADERS[f] for f in fields])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(vertical="center")

    amount_col = fields.index("amount") + 1
    for row_idx, tx in enumerate(transactions, start=2):
        for col_idx, field in enumerate(fields, start=1):
            value = _amount_value(tx.get(field)) if field == "amount" else str(tx.get(field) or "")
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if isinstance(value, str) and value.startswith("="):
                # openpyxl treats a leading "=" as a formula; keep it as plain text.
                cell.data_type = "s"
            if col_idx == amount_col and isinstance(value, float):
                cell.number_format = "#,##0.00"

    for col_idx, field in enumerate(fields, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = _WIDTHS[field]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()

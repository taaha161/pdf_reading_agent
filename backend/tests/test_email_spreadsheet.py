"""Tests for /api/email-spreadsheet (anonymous "email me the spreadsheet")."""
import io
from unittest.mock import patch

import pytest
from openpyxl import load_workbook

from main import _sanitize_rows

JOB_ID = "00000000-0000-0000-0000-000000000042"
CSV = "date,description,amount,type,category\n2024-01-01,Coffee,-4.50,debit,Food\n"
ROWS = [{"date": "2024-01-01", "description": "Coffee", "amount": "-4.50", "type": "debit", "category": "Food"}]


@pytest.fixture
def deps():
    with patch("main.trial_run_exists", return_value=True) as exists, patch(
        "main.count_email_captures_for_job", return_value=0
    ) as count, patch("main.send_spreadsheet_email", return_value=True) as send, patch(
        "main.record_email_capture"
    ) as record, patch("main.notify_email_capture"):
        yield {"exists": exists, "count": count, "send": send, "record": record}


def _post(client, **overrides):
    payload = {"job_id": JOB_ID, "email": "Jane@Example.com", "transactions": ROWS, "marketing_opt_in": True}
    payload.update(overrides)
    return client.post("/api/email-spreadsheet", json=payload)


def test_sends_and_records_email(client, deps):
    r = _post(client)
    assert r.status_code == 200
    to, content, filename = deps["send"].call_args.args
    assert to == "jane@example.com"
    assert filename == "statement.csv"
    assert b"Coffee" in content and b"category" in content
    deps["record"].assert_called_once()
    assert deps["record"].call_args.args[:3] == ("jane@example.com", JOB_ID, True)


def test_sends_uncategorized_excel(client, deps):
    assert _post(client, format="xlsx", include_categories=False).status_code == 200
    _, content, filename = deps["send"].call_args.args
    assert filename == "statement-uncategorized.xlsx"
    ws = load_workbook(io.BytesIO(content)).active
    assert [c.value for c in ws[1]] == ["Date", "Description", "Amount", "Type"]
    assert ws["C2"].value == -4.5


def test_legacy_csv_content_still_accepted(client, deps):
    r = client.post("/api/email-spreadsheet", json={"job_id": JOB_ID, "email": "a@b.co", "csv_content": CSV})
    assert r.status_code == 200
    _, content, filename = deps["send"].call_args.args
    assert filename == "statement.csv" and b"Coffee" in content


def test_unknown_job_is_rejected(client, deps):
    deps["exists"].return_value = False
    assert _post(client).status_code == 404
    deps["send"].assert_not_called()


def test_per_job_send_limit(client, deps):
    deps["count"].return_value = 4
    assert _post(client).status_code == 429
    deps["send"].assert_not_called()


def test_invalid_email_rejected(client, deps):
    assert _post(client, email="not-an-email").status_code == 422


def test_wrong_csv_header_rejected(client, deps):
    r = client.post("/api/email-spreadsheet", json={"job_id": JOB_ID, "email": "a@b.co", "csv_content": "hello,world\n1,2\n"})
    assert r.status_code == 400
    deps["send"].assert_not_called()


def test_missing_data_rejected(client, deps):
    r = client.post("/api/email-spreadsheet", json={"job_id": JOB_ID, "email": "a@b.co"})
    assert r.status_code == 400


def test_send_failure_does_not_record(client, deps):
    deps["send"].return_value = False
    assert _post(client).status_code == 502
    deps["record"].assert_not_called()


def test_sanitize_neutralizes_formulas_but_keeps_negative_amounts():
    out = _sanitize_rows([
        {"date": "2024-01-01", "description": '=HYPERLINK("http://evil")', "amount": "-12.00", "type": "debit", "category": "@SUM(1)"}
    ])[0]
    assert out["description"].startswith("'=HYPERLINK")
    assert out["category"] == "'@SUM(1)"
    assert out["amount"] == "-12.00"

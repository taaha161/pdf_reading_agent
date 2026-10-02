"""Tests for /api/export (CSV / Excel, with or without categories)."""
import io

from openpyxl import load_workbook

ROWS = [
    {"date": "2024-01-01", "description": "Coffee", "amount": "-4.50", "type": "debit", "category": "Food"},
    {"date": "2024-01-02", "description": "=1+1", "amount": "1,200.00", "type": "credit", "category": None},
]


def _export(client, **body):
    return client.post("/api/export", json={"transactions": ROWS, **body})


def test_csv_with_categories(client):
    r = _export(client, format="csv", include_categories=True)
    assert r.status_code == 200
    assert 'filename="statement.csv"' in r.headers["content-disposition"]
    assert r.text.splitlines()[0] == "date,description,amount,type,category"
    assert "Coffee,-4.50,debit,Food" in r.text


def test_csv_without_categories(client):
    r = _export(client, format="csv", include_categories=False)
    assert 'filename="statement-uncategorized.csv"' in r.headers["content-disposition"]
    assert r.text.splitlines()[0] == "date,description,amount,type"
    assert "Food" not in r.text


def test_xlsx_with_categories(client):
    r = _export(client, format="xlsx", include_categories=True)
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[1]] == ["Date", "Description", "Amount", "Type", "Category"]
    assert ws["C2"].value == -4.5 and ws["C3"].value == 1200.0
    assert ws["B3"].value == "=1+1" and ws["B3"].data_type == "s"  # text, not a formula
    assert ws["E3"].value in (None, "")


def test_xlsx_without_categories(client):
    r = _export(client, format="xlsx", include_categories=False)
    assert 'filename="statement-uncategorized.xlsx"' in r.headers["content-disposition"]
    ws = load_workbook(io.BytesIO(r.content)).active
    assert ws.max_column == 4


def test_rejects_unknown_format(client):
    assert _export(client, format="pdf").status_code == 422
